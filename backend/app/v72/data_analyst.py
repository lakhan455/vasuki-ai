from __future__ import annotations

import ast
import csv
import io
import json
import math
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from typing import Any

from app.services.chat_v10 import route_chat_v10
from app.v12.sandbox import run_sandbox, sandbox_status

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 120
PREVIEW_ROWS = 12

_ALLOWED_PYTHON_IMPORTS = {
    "csv",
    "json",
    "math",
    "statistics",
    "collections",
    "datetime",
    "itertools",
    "re",
}
_FORBIDDEN_NAMES = {
    "__import__",
    "eval",
    "exec",
    "compile",
    "globals",
    "locals",
    "vars",
    "input",
    "breakpoint",
    "getattr",
    "setattr",
    "delattr",
}


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value
    text = str(value).strip()
    return text if text else None


def _as_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    if text.endswith("%"):
        try:
            number = float(text[:-1]) / 100.0
            return number if math.isfinite(number) else None
        except ValueError:
            return None
    for symbol in ("₹", "$", "€", "£"):
        text = text.replace(symbol, "")
    try:
        number = float(text)
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def _normalise_headers(values: list[Any]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for index, raw in enumerate(values[:MAX_COLUMNS]):
        base = str(raw or "").strip() or f"column_{index + 1}"
        seen[base] += 1
        headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return headers


def _records_from_matrix(matrix: list[list[Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    if not matrix:
        return [], []
    headers = _normalise_headers(matrix[0])
    records: list[dict[str, Any]] = []
    for raw in matrix[1 : MAX_ROWS + 1]:
        values = list(raw[: len(headers)]) + [None] * max(0, len(headers) - len(raw))
        records.append({key: _clean_cell(value) for key, value in zip(headers, values)})
    return headers, records


def _parse_csv(data: bytes) -> tuple[list[str], list[dict[str, Any]], str]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    headers, records = _records_from_matrix(rows)
    return headers, records, "csv"


def _parse_json(data: bytes) -> tuple[list[str], list[dict[str, Any]], str]:
    parsed = json.loads(data.decode("utf-8-sig", errors="strict"))
    if isinstance(parsed, dict):
        for key in ("rows", "data", "items", "records"):
            if isinstance(parsed.get(key), list):
                parsed = parsed[key]
                break
        else:
            parsed = [parsed]
    if not isinstance(parsed, list):
        raise ValueError("JSON must contain an object or an array of records.")
    objects = [item for item in parsed[:MAX_ROWS] if isinstance(item, dict)]
    headers: list[str] = []
    for item in objects:
        for key in item:
            name = str(key)
            if name not in headers:
                headers.append(name)
            if len(headers) >= MAX_COLUMNS:
                break
        if len(headers) >= MAX_COLUMNS:
            break
    records = [
        {key: _clean_cell(item.get(key)) for key in headers}
        for item in objects
    ]
    return headers, records, "json"


def _parse_xlsx(data: bytes) -> tuple[list[str], list[dict[str, Any]], str]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("XLSX analysis requires openpyxl.") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    matrix: list[list[Any]] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        if index > MAX_ROWS:
            break
        matrix.append(list(row[:MAX_COLUMNS]))
    headers, records = _records_from_matrix(matrix)
    return headers, records, "xlsx"


def load_table(filename: str, data: bytes) -> tuple[list[str], list[dict[str, Any]], str]:
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Data file exceeds the 15 MB V72 limit.")
    lower = (filename or "").casefold()
    if lower.endswith((".csv", ".tsv")):
        return _parse_csv(data)
    if lower.endswith(".json"):
        return _parse_json(data)
    if lower.endswith(".xlsx"):
        return _parse_xlsx(data)
    raise ValueError("V72 supports .csv, .tsv, .json and .xlsx files.")


def _quantile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    if not ordered:
        return 0.0
    if len(ordered) == 1:
        return ordered[0]
    position = (len(ordered) - 1) * fraction
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    weight = position - lower
    return ordered[lower] * (1.0 - weight) + ordered[upper] * weight


def _numeric_profile(name: str, values: list[float]) -> dict[str, Any]:
    q1 = _quantile(values, 0.25)
    q3 = _quantile(values, 0.75)
    iqr = q3 - q1
    low = q1 - 1.5 * iqr
    high = q3 + 1.5 * iqr
    outliers = sum(1 for value in values if value < low or value > high) if iqr > 0 else 0
    return {
        "name": name,
        "count": len(values),
        "sum": round(sum(values), 6),
        "min": round(min(values), 6),
        "max": round(max(values), 6),
        "mean": round(statistics.fmean(values), 6),
        "median": round(statistics.median(values), 6),
        "stdev": round(statistics.stdev(values), 6) if len(values) > 1 else 0.0,
        "q1": round(q1, 6),
        "q3": round(q3, 6),
        "outliers_iqr": outliers,
    }


def _column_profiles(headers: list[str], records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for name in headers:
        raw = [row.get(name) for row in records]
        present = [value for value in raw if value is not None and str(value).strip() != ""]
        numeric = [number for value in present if (number := _as_number(value)) is not None]
        unique_values = Counter(str(value) for value in present)
        profile: dict[str, Any] = {
            "name": name,
            "missing": len(raw) - len(present),
            "non_missing": len(present),
            "unique": len(unique_values),
            "numeric_ratio": round(len(numeric) / max(1, len(present)), 4),
        }
        if present and len(numeric) >= max(2, int(len(present) * 0.8)):
            profile["type"] = "numeric"
            profile["stats"] = _numeric_profile(name, numeric)
        else:
            profile["type"] = "categorical"
            profile["top_values"] = [
                {"value": value, "count": count}
                for value, count in unique_values.most_common(8)
            ]
        result.append(profile)
    return result


def _pearson(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < 3 or len(xs) != len(ys):
        return None
    mean_x = statistics.fmean(xs)
    mean_y = statistics.fmean(ys)
    dx = [x - mean_x for x in xs]
    dy = [y - mean_y for y in ys]
    denom = math.sqrt(sum(value * value for value in dx) * sum(value * value for value in dy))
    if denom == 0:
        return None
    return sum(a * b for a, b in zip(dx, dy)) / denom


def _correlations(headers: list[str], records: list[dict[str, Any]], profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    numeric_names = [item["name"] for item in profiles if item.get("type") == "numeric"][:12]
    pairs: list[dict[str, Any]] = []
    for index, left in enumerate(numeric_names):
        for right in numeric_names[index + 1 :]:
            xs: list[float] = []
            ys: list[float] = []
            for row in records:
                x = _as_number(row.get(left))
                y = _as_number(row.get(right))
                if x is not None and y is not None:
                    xs.append(x)
                    ys.append(y)
            value = _pearson(xs, ys)
            if value is not None:
                pairs.append({
                    "left": left,
                    "right": right,
                    "r": round(value, 4),
                    "pairs": len(xs),
                })
    pairs.sort(key=lambda item: abs(float(item["r"])), reverse=True)
    return pairs[:12]


def _group_insights(records: list[dict[str, Any]], profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    categorical = [
        item for item in profiles
        if item.get("type") == "categorical" and 1 < int(item.get("unique") or 0) <= 50
    ][:3]
    numeric = [item for item in profiles if item.get("type") == "numeric"][:3]
    insights: list[dict[str, Any]] = []
    for category in categorical:
        for measure in numeric:
            grouped: dict[str, list[float]] = defaultdict(list)
            for row in records:
                key = row.get(category["name"])
                number = _as_number(row.get(measure["name"]))
                if key is not None and number is not None:
                    grouped[str(key)].append(number)
            rows = [
                {
                    "group": key,
                    "count": len(values),
                    "sum": round(sum(values), 6),
                    "mean": round(statistics.fmean(values), 6),
                }
                for key, values in grouped.items()
                if values
            ]
            rows.sort(key=lambda item: abs(float(item["sum"])), reverse=True)
            if rows:
                insights.append({
                    "group_by": category["name"],
                    "measure": measure["name"],
                    "top_groups_by_sum": rows[:8],
                })
            if len(insights) >= 5:
                return insights
    return insights


def _parse_date(value: Any) -> datetime | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    candidates = [text, text.replace("Z", "+00:00")]
    for candidate in candidates:
        try:
            return datetime.fromisoformat(candidate)
        except ValueError:
            pass
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def _trend_insight(headers: list[str], records: list[dict[str, Any]], profiles: list[dict[str, Any]]) -> dict[str, Any] | None:
    numeric = [item["name"] for item in profiles if item.get("type") == "numeric"]
    if not numeric:
        return None
    date_name = ""
    parsed_dates: dict[int, datetime] = {}
    for name in headers:
        successes: dict[int, datetime] = {}
        for index, row in enumerate(records[:1000]):
            parsed = _parse_date(row.get(name))
            if parsed is not None:
                successes[index] = parsed
        if len(successes) >= max(3, int(min(len(records), 1000) * 0.6)):
            date_name = name
            parsed_dates = successes
            break
    if not date_name:
        return None
    measure = numeric[0]
    points: list[tuple[datetime, float]] = []
    for index, row in enumerate(records):
        date = parsed_dates.get(index) if index in parsed_dates else _parse_date(row.get(date_name))
        number = _as_number(row.get(measure))
        if date is not None and number is not None:
            points.append((date, number))
    if len(points) < 3:
        return None
    points.sort(key=lambda item: item[0])
    first = points[0]
    last = points[-1]
    change = last[1] - first[1]
    pct = (change / first[1] * 100.0) if first[1] != 0 else None
    return {
        "date_column": date_name,
        "measure": measure,
        "first": {"date": first[0].isoformat(), "value": round(first[1], 6)},
        "last": {"date": last[0].isoformat(), "value": round(last[1], 6)},
        "absolute_change": round(change, 6),
        "percent_change": round(pct, 4) if pct is not None else None,
        "points": len(points),
    }


def analyze_records(filename: str, data: bytes) -> dict[str, Any]:
    headers, records, kind = load_table(filename, data)
    if not headers:
        raise ValueError("The data file has no usable columns.")
    profiles = _column_profiles(headers, records)
    row_keys = [
        json.dumps(row, ensure_ascii=False, sort_keys=True, default=str)
        for row in records
    ]
    duplicates = len(row_keys) - len(set(row_keys))
    missing_cells = sum(int(item.get("missing") or 0) for item in profiles)
    return {
        "file": filename,
        "format": kind,
        "rows": len(records),
        "columns_count": len(headers),
        "headers": headers,
        "missing_cells": missing_cells,
        "duplicate_rows": duplicates,
        "columns": profiles,
        "correlations": _correlations(headers, records, profiles),
        "group_insights": _group_insights(records, profiles),
        "trend": _trend_insight(headers, records, profiles),
        "preview": records[:PREVIEW_ROWS],
        "limits": {
            "max_bytes": MAX_FILE_BYTES,
            "max_rows": MAX_ROWS,
            "max_columns": MAX_COLUMNS,
        },
    }


def _fallback_answer(report: dict[str, Any]) -> str:
    rows = int(report.get("rows") or 0)
    columns = int(report.get("columns_count") or 0)
    missing = int(report.get("missing_cells") or 0)
    duplicates = int(report.get("duplicate_rows") or 0)
    lines = [
        f"Analyzed {rows} rows across {columns} columns.",
        f"Missing cells: {missing}. Duplicate rows: {duplicates}.",
    ]
    correlations = report.get("correlations") or []
    if correlations:
        best = correlations[0]
        lines.append(
            f"Strongest numeric relationship found: {best['left']} vs {best['right']} (r={best['r']})."
        )
    trend = report.get("trend")
    if isinstance(trend, dict):
        lines.append(
            f"Detected time trend for {trend['measure']} using {trend['date_column']}: change {trend['absolute_change']}."
        )
    return "\n\n".join(lines)


async def analyze_data_file(
    settings: Any,
    *,
    filename: str,
    data: bytes,
    prompt: str,
) -> dict[str, Any]:
    report = analyze_records(filename, data)
    compact = {
        key: report[key]
        for key in (
            "file",
            "format",
            "rows",
            "columns_count",
            "missing_cells",
            "duplicate_rows",
            "columns",
            "correlations",
            "group_insights",
            "trend",
            "preview",
        )
    }
    instruction = f"""You are Vasuki V72 Advanced Data Analyst.
Answer the user's request using ONLY the computed dataset evidence below.
Do not invent rows, totals, trends, causes, or statistical significance.
If the evidence cannot answer something, say that clearly.
Prefer exact values and concise business-friendly interpretation.

USER REQUEST:
{str(prompt or '').strip() or 'Analyze this dataset and highlight the most important findings.'}

COMPUTED DATASET EVIDENCE:
{json.dumps(compact, ensure_ascii=False, default=str)[:100000]}
"""
    provider = "vasuki-v72-deterministic"
    answer = ""
    try:
        answer, provider = await route_chat_v10(
            "auto",
            [{"role": "user", "content": instruction}],
            settings,
            json.dumps(compact, ensure_ascii=False, default=str)[:80000],
        )
    except Exception:
        answer = ""
    if not str(answer or "").strip():
        answer = _fallback_answer(report)

    sandbox = sandbox_status()
    return {
        "ok": True,
        "version": "v72",
        "answer": str(answer).strip(),
        "provider": provider,
        "analysis": report,
        "execution_proof": {
            "engine": "v72-tabular-analysis",
            "dataset_rows_profiled": report["rows"],
            "dataset_columns_profiled": report["columns_count"],
            "computed_statistics": True,
            "correlation_analysis": True,
            "group_aggregation": True,
            "trend_detection": report.get("trend") is not None,
            "python_sandbox_used": False,
            "bounded_python_sandbox_available": bool(sandbox.get("available")),
            "network_used_for_dataset_computation": False,
        },
    }


def _normalised_csv_bytes(filename: str, data: bytes) -> bytes:
    headers, records, _kind = load_table(filename, data)
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    for row in records:
        writer.writerow({key: "" if row.get(key) is None else row.get(key) for key in headers})
    return output.getvalue().encode("utf-8")


def validate_bounded_python(code: str) -> None:
    if not code.strip():
        raise ValueError("Python code is empty.")
    if len(code) > 20_000:
        raise ValueError("Python analysis code exceeds 20,000 characters.")
    try:
        tree = ast.parse(code)
    except SyntaxError as exc:
        raise ValueError(f"Python syntax error at line {exc.lineno}: {exc.msg}") from exc
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                root = alias.name.split(".", 1)[0]
                if root not in _ALLOWED_PYTHON_IMPORTS:
                    raise ValueError(f"Import not allowed in V72 sandbox: {alias.name}")
        elif isinstance(node, ast.ImportFrom):
            root = (node.module or "").split(".", 1)[0]
            if root not in _ALLOWED_PYTHON_IMPORTS:
                raise ValueError(f"Import not allowed in V72 sandbox: {node.module}")
        elif isinstance(node, ast.Name) and node.id in _FORBIDDEN_NAMES:
            raise ValueError(f"Name not allowed in V72 sandbox: {node.id}")
        elif isinstance(node, ast.Attribute) and node.attr.startswith("__"):
            raise ValueError("Dunder attribute access is not allowed in V72 sandbox.")


async def execute_bounded_python(
    *,
    filename: str,
    data: bytes,
    code: str,
) -> dict[str, Any]:
    validate_bounded_python(code)
    status = sandbox_status()
    if not status.get("available"):
        return {
            "ok": False,
            "available": False,
            "version": "v72",
            "error": "Bounded Python execution requires the existing Docker sandbox, which is unavailable on this runtime.",
            "sandbox": status,
            "arbitrary_server_python_execution": False,
        }
    csv_bytes = _normalised_csv_bytes(filename, data)
    files = {
        "analysis.py": code,
        "data.csv": csv_bytes.decode("utf-8", errors="replace"),
    }
    result = await run_sandbox(
        files,
        runtime="python",
        command="python analysis.py",
        timeout_seconds=25,
    )
    return {
        "version": "v72",
        "bounded": True,
        "allowed_imports": sorted(_ALLOWED_PYTHON_IMPORTS),
        "data_filename_inside_sandbox": "data.csv",
        **result,
    }


def health_snapshot() -> dict[str, Any]:
    sandbox = sandbox_status()
    return {
        "ok": True,
        "version": "v72",
        "name": "Vasuki Advanced Data Analyst",
        "supported_files": ["csv", "tsv", "json", "xlsx"],
        "features": [
            "full-dataset-column-profiling",
            "missing-and-duplicate-detection",
            "numeric-summary-statistics",
            "iqr-outlier-counts",
            "pearson-correlation-ranking",
            "categorical-group-aggregations",
            "basic-date-trend-detection",
            "evidence-grounded-model-summary",
            "bounded-python-when-docker-is-available",
        ],
        "deterministic_data_computation": True,
        "bounded_python_execution_available": bool(sandbox.get("available")),
        "arbitrary_server_python_execution": False,
        "sandbox": sandbox,
        "new_database_migration_required": False,
        "new_api_key_required": False,
    }
